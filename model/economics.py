"""Экономическая модель AI-маркетолога для Города.

Запуск:  python3 model/economics.py           — печатает сводные таблицы
         python3 model/economics.py --xlsx    — ещё и собирает model/economics.xlsx (нужен openpyxl)

Источники входов:
  B     — базовая линия (docs/00-baseline.md), условия кейса
  A#    — допущения из docs/assumptions.md
  P1-#  — docs/01-product.md, P2-# — docs/02-autonomy.md, PA-# — docs/03-architecture.md
  P3-#  — docs/06-economics-pilot.md
Деньги — млн ₽ в месяц, если не сказано иное. Дни — рабочие.
"""
import math
import sys

# ============================================================================
# 1. ВХОДЫ
# ============================================================================

BASE = {
    "campaigns_per_month": 120,                            # B
    "fte": {"marketer": 14, "analyst": 6, "editor": 4},    # B, A1 (без проджектов)
    "fte_cost": 0.35,                                      # B
    "comms_budget": 40.0,                                  # B
    "launch_days_asis": 9.0,                               # B, A3
    "debrief_asis": 5.0,                                   # B
}

# Сроки по классам риска: дни до отправки по классу (P1-6) и доли классов (P2-4)
CLASS_DAYS = {"К0": 1.0, "К1": 1.0, "К2": 2.0, "К3": 3.5}  # P1-6

SCEN_NAMES = ("конс", "база", "опт")
# поле: (ID, конс, база, опт)
S = {
    # --- FTE: рутина, которую забирает агент, и новая работа, по ролям ---
    "routine_marketer": ("P1-3 / P3-1", 0.45, 0.55, 0.60),
    "routine_analyst":  ("P1-3 / P3-1", 0.55, 0.65, 0.70),
    "routine_editor":   ("P1-3 / P3-1", 0.50, 0.60, 0.70),
    "newwork_marketer": ("P1-3 / P3-2", 0.20, 0.15, 0.15),   # утверждение, каталог условий, бэклог
    "newwork_analyst":  ("P1-3 / P3-2", 0.25, 0.20, 0.20),   # методика, метрики агента и судьи
    "newwork_editor":   ("P1-3 / P3-2", 0.25, 0.20, 0.20),   # гайд тона, эталоны, разметка, аудит К0
    "realization":      ("P3-3", 0.60, 0.80, 0.90),         # доля времени, складывающаяся в целых людей
    # --- Сроки ---
    "class_mix_start":  ("P2-4 / P3-4", None, (0.00, 0.45, 0.35, 0.20), (0.00, 0.55, 0.30, 0.15)),
    "class_mix_6m":     ("P2-4 / P3-4", None, (0.40, 0.30, 0.20, 0.10), (0.55, 0.30, 0.10, 0.05)),
    "launch_fixed":     ("P1-1", 6.0, None, None),             # конс: нет API платформы рассылок
    "debrief_tobe":     ("P3-4", 2.0, 1.0, 1.0),
    # --- Коммуникации ---
    "comms_saving_share": ("P3-5", 0.00, 0.02, 0.05),
    # --- Стоимость системы (03-architecture.md) ---
    "dev_fte_cost":     ("P3-6", 0.75, 0.60, 0.50),          # млн ₽/мес на инженера с налогами
    "pilot_team":       ("PA-8", 6, 6, 6),                  # 5 + интеграционный бэкенд с м1
    "support_team":     ("PA-6", 3.0, 2.0, 1.5),
    "infra_month":      ("PA-5", 0.6, 0.25, 0.10),
    "inference_per_campaign_rub": ("PA-4", 2_000, 500, 150),
    "one_off_setup":    ("PA-6", 2.5, 1.2, 0.6),
    # --- Внедрение ---
    "rollout_months":   ("P3-11", 9, 6, 4),
    "release_lag":      ("P3-11", 3, 2, 1),
    # --- Апсайд качества (вне окупаемости, A5 открыт) ---
    "uplift_rel_change": ("P3-12", 0.00, 0.05, 0.15),
}
PILOT_MONTHS = 6
PILOT_AGENT_CAMPAIGNS = [0, 20, 15, 25, 25, 30]   # P3-13 (м2 — теневой режим; м4–м6 — 80 рандомизированных)
HORIZON = 36
INCR_VALUE_FLOOR = 40.0   # P3-12: кампании как минимум окупают свои коммуникации
INFERENCE_STOP_RUB = 5_000  # стоп-порог инференса на кампанию (06-economics-pilot.md, раздел 4)


def scen(i):
    return {k: v[1 + i] for k, v in S.items()}


# ============================================================================
# 2. РАСЧЁТ
# ============================================================================

ROLES = (("marketer", "Маркетологи"), ("analyst", "Аналитики"), ("editor", "Редакторы"))


def fte_block(p):
    rows, routine, newwork = {}, 0.0, 0.0
    for role, _ in ROLES:
        n = BASE["fte"][role]
        r, w = p[f"routine_{role}"], p[f"newwork_{role}"]
        rows[role] = (n, r, w, n * (r - w))
        routine += n * r
        newwork += n * w
    net = routine - newwork
    released = net * p["realization"]
    return {"rows": rows, "routine": routine, "newwork": newwork, "net": net,
            "released": released, "money": released * BASE["fte_cost"]}


def mix_days(mix):
    return sum(s * d for s, d in zip(mix, CLASS_DAYS.values()))


def days_block(p):
    if p["launch_fixed"] is not None:
        start = mature = p["launch_fixed"]
    else:
        start, mature = mix_days(p["class_mix_start"]), mix_days(p["class_mix_6m"])
    return {"start": start, "mature": mature, "saved_mature": BASE["launch_days_asis"] - mature,
            "debrief": p["debrief_tobe"], "debrief_saved": BASE["debrief_asis"] - p["debrief_tobe"]}


def coverage_curve(p):
    cov = [0.0] * HORIZON
    for t in range(PILOT_MONTHS, HORIZON):
        cov[t] = min(1.0, (t - PILOT_MONTHS + 1) / p["rollout_months"])
    return cov


def quality_upside(p):
    return p["uplift_rel_change"] * INCR_VALUE_FLOOR


def cashflow(p, with_comms=False, extra=0.0, with_quality=False):
    """Помесячный поток. Пилот: экономию не считаем. Тираж: команда пилота остаётся до 100%
    покрытия, затем поддержка. FTE высвобождаются с лагом release_lag."""
    fte = fte_block(p)
    cov = coverage_curve(p)
    rows, cum, payback = [], 0.0, None
    for t in range(HORIZON):
        in_pilot = t < PILOT_MONTHS
        n_agent = PILOT_AGENT_CAMPAIGNS[t] if in_pilot else BASE["campaigns_per_month"] * cov[t]
        inference = n_agent * p["inference_per_campaign_rub"] / 1e6
        team_size = p["pilot_team"] if (in_pilot or cov[t] < 1.0) else p["support_team"]
        team = team_size * p["dev_fte_cost"]
        one_off = p["one_off_setup"] if t == 0 else 0.0
        cost = team + p["infra_month"] + inference + one_off
        lag_t = t - int(p["release_lag"])
        save_fte = fte["money"] * (cov[lag_t] if lag_t >= 0 else 0.0)
        save_comms = BASE["comms_budget"] * p["comms_saving_share"] * cov[t] if with_comms else 0.0
        save_q = quality_upside(p) * cov[t] if with_quality else 0.0
        net = save_fte + save_comms + save_q + extra * cov[t] - cost
        cum += net
        if payback is None and cum >= 0 and t >= PILOT_MONTHS:
            payback = t + 1
        rows.append({"m": t + 1, "cov": cov[t], "team": team, "infra": p["infra_month"],
                     "inference": inference, "one_off": one_off, "cost": cost,
                     "save_fte": save_fte, "save_comms": save_comms, "net": net, "cum": cum})
    return rows, payback


def steady_state(p):
    fte = fte_block(p)
    run = p["support_team"] * p["dev_fte_cost"] + p["infra_month"] + \
        BASE["campaigns_per_month"] * p["inference_per_campaign_rub"] / 1e6
    comms = BASE["comms_budget"] * p["comms_saving_share"]
    return {"save_fte": fte["money"], "save_comms": comms, "run": run,
            "net_fte": fte["money"] - run, "net_all": fte["money"] + comms - run}


def extra_needed(p, target_month):
    """Дополнительная выгода сверх FTE (млн ₽/мес при 100% покрытии), нужная для окупаемости
    к target_month. Бинарный поиск."""
    lo, hi = 0.0, 100.0
    for _ in range(60):
        mid = (lo + hi) / 2
        _, pb = cashflow(p, extra=mid)
        if pb is not None and pb <= target_month:
            hi = mid
        else:
            lo = mid
    return hi


SENS_GROUPS = {
    "Рутина, которую забирает агент (P1-3/P3-1)": ["routine_marketer", "routine_analyst", "routine_editor"],
    "Новая работа людей (P1-3/P3-2)": ["newwork_marketer", "newwork_analyst", "newwork_editor"],
    "Реализуемость в целых людях (P3-3)": ["realization"],
    "Стоимость FTE инженера (P3-6)": ["dev_fte_cost"],
    "Команда поддержки (PA-6)": ["support_team"],
    "Скорость тиража и лаг (P3-11)": ["rollout_months", "release_lag"],
    "Инфраструктура (PA-5)": ["infra_month"],
    "Инференс на кампанию (PA-4)": ["inference_per_campaign_rub"],
    "Разовые затраты (PA-6)": ["one_off_setup"],
}


def sensitivity():
    base = scen(1)
    rows_b, pb_b = cashflow(base)
    out = []
    for name, keys in SENS_GROUPS.items():
        res = []
        for i in (0, 2):
            p = dict(base)
            for k in keys:
                p[k] = S[k][1 + i]
            rows, pb = cashflow(p)
            res.append((rows[-1]["cum"], pb))
        out.append((name, res[0], res[1], abs(res[1][0] - res[0][0])))
    out.sort(key=lambda x: -x[3])
    return rows_b[-1]["cum"], pb_b, out


def power_n(sigma_pp, delta_pp):
    """Кампаний на плечо для «не хуже» (α = 5% одностороннее, мощность 80%)."""
    return math.ceil(2 * (1.6449 + 0.8416) ** 2 * sigma_pp ** 2 / delta_pp ** 2)


# ============================================================================
# 3. ПЕЧАТЬ
# ============================================================================

def fmt(x, d=1):
    if x is None:
        return f">{HORIZON}"
    if isinstance(x, str):
        return x
    if isinstance(x, int):
        return str(x)
    return f"{x:,.{d}f}".replace(",", " ").replace(".", ",")


def table(headers, rows):
    w = [max(len(str(h)), *(len(str(r[i])) for r in rows)) for i, h in enumerate(headers)]
    print(" | ".join(str(h).ljust(w[i]) for i, h in enumerate(headers)))
    print("-+-".join("-" * x for x in w))
    for r in rows:
        print(" | ".join(str(c).ljust(w[i]) for i, c in enumerate(r)))
    print()


def head(t):
    print("=" * 78)
    print(t)
    print("=" * 78)


def collect():
    R = {}
    for i, n in enumerate(SCEN_NAMES):
        p = scen(i)
        cf, pb = cashflow(p)
        cf2, pb2 = cashflow(p, with_comms=True)
        R[n] = {"p": p, "fte": fte_block(p), "days": days_block(p), "cf": cf, "pb": pb,
                "cf2": cf2, "pb2": pb2, "ss": steady_state(p), "upside": quality_upside(p)}
    return R


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")   # консоль Windows по умолчанию не UTF-8
    R = collect()
    N = SCEN_NAMES

    head("1. ВЫСВОБОЖДЕНИЕ FTE (база 24 FTE, A1; доли P1-3)")
    rows = []
    for role, name in ROLES:
        rows.append([name, BASE["fte"][role]] +
                    [f"{R[n]['fte']['rows'][role][1]:.0%} − {R[n]['fte']['rows'][role][2]:.0%} → "
                     f"{fmt(R[n]['fte']['rows'][role][3])}" for n in N])
    rows.append(["Рутина, которую забирает агент", 24] + [fmt(R[n]["fte"]["routine"]) for n in N])
    rows.append(["− новая работа", ""] + [fmt(R[n]["fte"]["newwork"]) for n in N])
    rows.append(["= чистое высвобождение времени", ""] + [fmt(R[n]["fte"]["net"]) for n in N])
    rows.append(["× реализуемость (P3-3)", ""] + [f"{R[n]['p']['realization']:.0%}" for n in N])
    rows.append(["ИТОГО высвобождено FTE", ""] + [fmt(R[n]["fte"]["released"]) for n in N])
    rows.append(["Доля от 24 FTE", ""] + [f"{R[n]['fte']['released'] / 24:.0%}" for n in N])
    rows.append(["Экономия, млн ₽/мес", ""] + [fmt(R[n]["fte"]["money"], 2) for n in N])
    table(["Роль (рутина − новая работа)", "FTE"] + list(N), rows)

    head("2. СРОКИ (рабочие дни; база по классам К0–К3, P1-6 × P2-4)")
    rows = [["До запуска, старт пилота", fmt(BASE["launch_days_asis"])] + [fmt(R[n]["days"]["start"], 2) for n in N],
            ["До запуска, через 6 мес", fmt(BASE["launch_days_asis"])] + [fmt(R[n]["days"]["mature"], 2) for n in N],
            ["Экономия дней запуска (зрелость)", ""] + [fmt(R[n]["days"]["saved_mature"], 2) for n in N],
            ["Разбор итогов", fmt(BASE["debrief_asis"])] + [fmt(R[n]["days"]["debrief"]) for n in N]]
    table(["", "как есть"] + list(N), rows)

    head("3. СТОИМОСТЬ СИСТЕМЫ (03-architecture.md)")
    rows = []
    for n in N:
        p, cf = R[n]["p"], R[n]["cf"]
        pilot = sum(r["cost"] for r in cf[:PILOT_MONTHS])
        team = sum(r["team"] for r in cf[:PILOT_MONTHS])
        rows.append([n, fmt(team, 1), fmt(p["infra_month"] * PILOT_MONTHS, 2),
                     fmt(sum(r["inference"] for r in cf[:PILOT_MONTHS]), 2), fmt(p["one_off_setup"], 1),
                     fmt(pilot, 1), fmt(R[n]["ss"]["run"], 2)])
    table(["Сценарий", "Команда 6 × 6 мес", "Инфра за пилот", "Инференс за пилот", "Разовые",
           "ПИЛОТ всего", "Эксплуатация /мес"], rows)

    head("4. УСТАНОВИВШИЙСЯ МЕСЯЦ (100% покрытие, FTE высвобождены), млн ₽/мес")
    rows = [[n, fmt(R[n]["ss"]["save_fte"], 2), fmt(R[n]["ss"]["run"], 2), fmt(R[n]["ss"]["net_fte"], 2),
             fmt(R[n]["ss"]["save_comms"], 2), fmt(R[n]["ss"]["net_all"], 2), fmt(R[n]["upside"], 1)]
            for n in N]
    table(["Сценарий", "Экономия FTE", "Эксплуатация", "Чистый: FTE", "Каналы (P3-5)",
           "Чистый: FTE+каналы", "Апсайд качества*"], rows)
    print(f"* вне окупаемости: нет среднего чека (A5); ценность инкремента ≥ {fmt(INCR_VALUE_FLOOR, 0)} млн ₽/мес (P3-12)\n")

    head(f"5. ОКУПАЕМОСТЬ (месяц, когда накопленный итог ≥ 0; горизонт {HORIZON} мес)")
    rows = [[n, fmt(R[n]["pb"]), fmt(R[n]["cf"][11]["cum"]), fmt(R[n]["cf"][23]["cum"]), fmt(R[n]["cf"][35]["cum"]),
             fmt(R[n]["pb2"]), fmt(R[n]["cf2"][35]["cum"])] for n in N]
    table(["Сценарий", "Окуп.: FTE", "Нак. 12", "Нак. 24", "Нак. 36", "Окуп.: FTE+каналы", "Нак. 36: FTE+каналы"], rows)

    print("Пик вложений: минимум накопленного итога (пилот + тираж до выхода экономии)")
    rows = []
    for n in N:
        lo = min(R[n]["cf"], key=lambda r: r["cum"])
        lo2 = min(R[n]["cf2"], key=lambda r: r["cum"])
        rows.append([n, fmt(lo["cum"]), lo["m"], fmt(lo2["cum"]), lo2["m"]])
    table(["Сценарий", "Пик: FTE", "Месяц", "Пик: FTE+каналы", "Месяц"], rows)

    print("Порог окупаемости: сколько нужно сверх FTE (при 100% покрытии)")
    rows = []
    for n in N:
        for tm in (18, 24):
            x = extra_needed(R[n]["p"], tm)
            rows.append([n, tm, fmt(x, 2), f"{fmt(100 * x / BASE['comms_budget'])}%",
                         f"+{fmt(100 * x / INCR_VALUE_FLOOR)}%"])
    table(["Сценарий", "К месяцу", "млн ₽/мес", "% бюджета коммуникаций", "к инкрем. эффекту (P3-12)"], rows)
    _, pb3 = cashflow(R["база"]["p"], with_comms=True, with_quality=True)
    print(f"Иллюстрация: база FTE + каналы 2% + качество +5% — окупаемость месяц {fmt(pb3)}.\n")

    print("Помесячно, базовый сценарий (только FTE), млн ₽:")
    rows = [[r["m"], f"{r['cov']:.0%}", fmt(r["cost"], 2), fmt(r["save_fte"], 2), fmt(r["net"], 2), fmt(r["cum"], 1)]
            for r in R["база"]["cf"] if r["m"] <= 15 or r["m"] % 6 == 0]
    table(["Мес", "Покрытие", "Затраты", "Экономия FTE", "Чистый", "Накопленный"], rows)

    head(f"6. ЧУВСТВИТЕЛЬНОСТЬ (база, только FTE; накопленный итог за {HORIZON} мес, млн ₽)")
    base_cum, pb_b, sens = sensitivity()
    print(f"База: накопленный итог {fmt(base_cum)}, окупаемость — месяц {fmt(pb_b)}\n")
    table(["Параметр (конс ↔ опт, остальное — база)", "Конс", "Опт", "Размах"],
          [[nm, f"{fmt(lo[0])} (м. {fmt(lo[1])})", f"{fmt(hi[0])} (м. {fmt(hi[1])})", fmt(sw)]
           for nm, lo, hi, sw in sens])

    head("7. МОЩНОСТЬ ПИЛОТА: кампаний на плечо для «агент не хуже людей» (P3-14)")
    table(["Разброс эффекта между кампаниями", "δ = 4 п.п.", "δ = 6 п.п.", "δ = 8 п.п."],
          [[f"σ = {s} п.п."] + [power_n(s, d) for d in (4, 6, 8)] for s in (10, 15, 20)])
    print(f"Стоп-порог инференса: {INFERENCE_STOP_RUB} ₽ на кампанию = "
          f"{INFERENCE_STOP_RUB / S['inference_per_campaign_rub'][2]:.0f}× базы.\n")

    # --- проверки арифметики ---
    b = R["база"]
    assert sum(BASE["fte"].values()) == 24 and abs(24 * BASE["fte_cost"] - 8.4) < 1e-9
    assert abs(b["fte"]["net"] - (14 * .40 + 6 * .45 + 4 * .40)) < 1e-9          # P1-3: 40/45/40%
    assert abs(b["fte"]["released"] - 7.92) < 1e-9 and abs(b["fte"]["money"] - 2.772) < 1e-9
    assert abs(b["days"]["start"] - 1.85) < 1e-9 and abs(b["days"]["mature"] - 1.45) < 1e-9
    pilot = sum(r["cost"] for r in b["cf"][:PILOT_MONTHS])
    assert abs(pilot - 24.4) < 0.1, pilot                                       # пилот ≈ 24,4
    assert abs(b["ss"]["run"] - 1.5) < 0.05, b["ss"]["run"]                     # эксплуатация ≈ 1,5
    assert abs(b["cf"][-1]["cum"] - sum(r["net"] for r in b["cf"])) < 1e-9
    for n in N:
        assert R[n]["days"]["mature"] >= 1.0                                    # P1-6: не меньше 1 дня
    print("Проверки арифметики пройдены.")

    if "--xlsx" in sys.argv:
        export_xlsx(R)


def export_xlsx(R):
    try:
        from openpyxl import Workbook
    except ImportError:
        print("openpyxl не установлен — xlsx не собран")
        return
    import os
    wb = Workbook()
    ws = wb.active
    ws.title = "Входы"
    ws.append(["Значения, не формулы. Источник истины — model/economics.py"])
    ws.append(["Параметр", "ID", "Конс", "База", "Опт"])
    for k, v in S.items():
        ws.append([k, v[0]] + [str(x) if isinstance(x, tuple) else x for x in v[1:]])
    w = wb.create_sheet("Сводка")
    w.append(["Сценарий", "FTE высвобождено", "Экономия FTE, млн/мес", "Дни до запуска (зрелость)",
              "Пилот, млн", "Эксплуатация, млн/мес", "Окуп. FTE, мес", "Окуп. FTE+каналы, мес",
              "Порог 18 мес, % коммуникаций", "Порог 24 мес, % коммуникаций"])
    for n in SCEN_NAMES:
        r = R[n]
        w.append([n, round(r["fte"]["released"], 2), round(r["fte"]["money"], 3), round(r["days"]["mature"], 2),
                  round(sum(x["cost"] for x in r["cf"][:PILOT_MONTHS]), 2), round(r["ss"]["run"], 3),
                  r["pb"] or f">{HORIZON}", r["pb2"] or f">{HORIZON}",
                  round(100 * extra_needed(r["p"], 18) / BASE["comms_budget"], 2),
                  round(100 * extra_needed(r["p"], 24) / BASE["comms_budget"], 2)])
    for n in SCEN_NAMES:
        w = wb.create_sheet(f"Поток {n}")
        w.append(["Месяц", "Покрытие", "Команда", "Инфра", "Инференс", "Разовые", "Затраты",
                  "Экономия FTE", "Чистый", "Накопленный"])
        for r in R[n]["cf"]:
            w.append([r["m"], round(r["cov"], 3), round(r["team"], 3), round(r["infra"], 3),
                      round(r["inference"], 4), r["one_off"], round(r["cost"], 3),
                      round(r["save_fte"], 3), round(r["net"], 3), round(r["cum"], 3)])
    w = wb.create_sheet("Чувствительность")
    _, _, sens = sensitivity()
    w.append(["Параметр", "Нак. 36 мес конс", "Окуп. конс", "Нак. 36 мес опт", "Окуп. опт", "Размах"])
    for nm, lo, hi, sw in sens:
        w.append([nm, round(lo[0], 2), lo[1] or f">{HORIZON}", round(hi[0], 2), hi[1] or f">{HORIZON}", round(sw, 2)])
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "economics.xlsx")
    wb.save(path)
    print(f"Сохранено: model/economics.xlsx")


if __name__ == "__main__":
    main()
